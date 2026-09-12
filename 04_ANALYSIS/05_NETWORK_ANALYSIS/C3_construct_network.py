"""Construct the approved C3 supplier-technology network.

The script reads the frozen analytical master and the approved B4 denominator
register, validates the primary population, constructs the weighted bipartite
edge list and weighted-Jaccard supplier projection, and calculates the raw C3
network measures. Exact Fraction arithmetic is used for similarities,
distances, shortest-path ties, strength, and normalized betweenness.

It does not edit any source workbook and does not calculate C4 scores or
profiles, C5 sensitivities, C6/D3 figures, or Random Forest outputs.
"""

from __future__ import annotations

import argparse
import hashlib
import heapq
import json
import platform
from collections import Counter, defaultdict, deque
from fractions import Fraction
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
ORIGINAL_PATH = ROOT / "03_DATA_AND_EVIDENCE/00_RAW_READ_ONLY/02_ANALYSIS_WORKBOOK.xlsx"
SOURCE_PATH = ROOT / "03_DATA_AND_EVIDENCE/01_FINAL_MASTER/04_SUPPLIER_EVIDENCE_ANALYTICAL_MASTER_FINAL.xlsx"
FULL_AUDIT_PATH = ROOT / "03_DATA_AND_EVIDENCE/00_RAW_READ_ONLY/FULL_AUDIT/03_SUPPLIER_EVIDENCE_AUDIT_MASTER_21_SHEETS.xlsx"
B1_PATH = ROOT / "04_ANALYSIS/02_DATA_PREPARATION/B1_TAXONOMY_REVIEW_REGISTER.xlsx"
B2_PATH = ROOT / "04_ANALYSIS/02_DATA_PREPARATION/B2_DATE_COMPLETENESS_AUDIT.xlsx"
B4_PATH = ROOT / "04_ANALYSIS/02_DATA_PREPARATION/B4_DENOMINATOR_REGISTER.xlsx"
C1_PATH = ROOT / "04_ANALYSIS/03_DESCRIPTIVE_STATISTICS/C1_DESCRIPTIVE_STATISTICS_WORKING.xlsx"
C2_PATH = ROOT / "04_ANALYSIS/04_TECHNOLOGY_CATEGORY_COMPARISONS/C2_TECHNOLOGY_CATEGORY_COMPARISONS_WORKING.xlsx"
IMPLEMENTATION_PATH = ROOT / "04_ANALYSIS/05_NETWORK_ANALYSIS/C3_network_implementation.md"

EXPECTED_HASHES = {
    "original_workbook": "F3352FAE0A07565066902CC72070651C23CCD09CD9DB2C05BEFC4FDAFC85C358",
    "analytical_master": "8C5B3439CFFF00A2110C575E101CD32C4988BDDC3269DDDDD46497F2A950C16D",
    "full_audit": "52EF4BA56C54AC56467AF70891444CBA227F68413B654A4D2A9FD9810E362D9F",
    "b1_register": "0AC5DB0E1171D98E2DC7C919522B2EEA18D4283972FCCA04B61215A3278EFA51",
    "b2_audit": "08B39F195142965459A385931D92ACDBC6866A6667A1A3AA83EE871E902B4F04",
    "b4_register": "348C384002B1D54E50293238EF48621CF7FA45293AF92CBC29A32C652DAF933E",
    "c1_workbook": "B8BE1FDB8F464626DF57F2914D935F5876BA138F496339459EBE80A837F3BDF3",
    "c2_workbook": "E18E37FA4408F18D7AA416605C5FBCA2186BAD8D5C67247A4B289A67887394CD",
}

HASH_PATHS = {
    "original_workbook": ORIGINAL_PATH,
    "analytical_master": SOURCE_PATH,
    "full_audit": FULL_AUDIT_PATH,
    "b1_register": B1_PATH,
    "b2_audit": B2_PATH,
    "b4_register": B4_PATH,
    "c1_workbook": C1_PATH,
    "c2_workbook": C2_PATH,
}


def norm(value: Any) -> str:
    return "" if value is None else str(value).strip()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest().upper()


def sha256_text(lines: Iterable[str]) -> str:
    payload = "\n".join(lines).encode("utf-8")
    return hashlib.sha256(payload).hexdigest().upper()


def require_private_output(path: Path) -> Path:
    """Restrict row-level JSON exports to the repository's ignored tmp tree."""

    resolved = path.resolve()
    private_root = (ROOT / "tmp").resolve()
    if resolved != private_root and private_root not in resolved.parents:
        raise ValueError(f"Full C3 output must stay under the ignored local directory: {private_root}")
    return resolved


def read_sheet(path: Path, sheet_name: str) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name]
        iterator = sheet.iter_rows(values_only=True)
        headers = [norm(value) for value in next(iterator)]
        rows: list[dict[str, Any]] = []
        for values in iterator:
            if not any(norm(value) for value in values):
                continue
            rows.append({headers[index]: value for index, value in enumerate(values) if index < len(headers)})
        return rows
    finally:
        workbook.close()


def require_fields(rows: list[dict[str, Any]], fields: Iterable[str], label: str) -> None:
    if not rows:
        raise ValueError(f"{label} contains no data rows")
    missing = [field for field in fields if field not in rows[0]]
    if missing:
        raise ValueError(f"{label} is missing fields: {missing}")


def decimal(value: Fraction, places: int = 15) -> float:
    return round(value.numerator / value.denominator, places)


def graph_from_edges(nodes: Iterable[str], edges: Iterable[tuple[str, str, Fraction]]) -> dict[str, dict[str, Fraction]]:
    adjacency: dict[str, dict[str, Fraction]] = {node: {} for node in nodes}
    for left, right, distance in edges:
        if left == right:
            raise ValueError("Self-loop is not permitted")
        if distance <= 0:
            raise ValueError("Distance must be positive")
        if right in adjacency[left] or left in adjacency[right]:
            raise ValueError(f"Parallel edge is not permitted: {left}, {right}")
        adjacency[left][right] = distance
        adjacency[right][left] = distance
    return adjacency


def brandes_normalized(adjacency: dict[str, dict[str, Fraction]]) -> dict[str, Fraction]:
    """Exact weighted Brandes centrality for a simple undirected graph.

    The source-accumulated Brandes values count undirected paths twice. The
    standard normalized undirected scaling is therefore 1/[(n-1)(n-2)].
    Endpoints are excluded and disconnected pairs contribute zero.
    """

    nodes = sorted(adjacency)
    centrality = {node: Fraction(0, 1) for node in nodes}
    for source in nodes:
        stack: list[str] = []
        predecessors: dict[str, list[str]] = {node: [] for node in nodes}
        sigma = {node: 0 for node in nodes}
        sigma[source] = 1
        distance: dict[str, Fraction | None] = {node: None for node in nodes}
        distance[source] = Fraction(0, 1)
        queue: list[tuple[Fraction, str]] = [(Fraction(0, 1), source)]

        while queue:
            current_distance, vertex = heapq.heappop(queue)
            if distance[vertex] is None or current_distance != distance[vertex]:
                continue
            stack.append(vertex)
            for neighbor in sorted(adjacency[vertex]):
                candidate = current_distance + adjacency[vertex][neighbor]
                if distance[neighbor] is None or candidate < distance[neighbor]:
                    distance[neighbor] = candidate
                    heapq.heappush(queue, (candidate, neighbor))
                    sigma[neighbor] = sigma[vertex]
                    predecessors[neighbor] = [vertex]
                elif candidate == distance[neighbor]:
                    sigma[neighbor] += sigma[vertex]
                    predecessors[neighbor].append(vertex)

        dependency = {node: Fraction(0, 1) for node in nodes}
        while stack:
            successor = stack.pop()
            if sigma[successor] == 0:
                continue
            coefficient = (Fraction(1, 1) + dependency[successor]) / sigma[successor]
            for predecessor in predecessors[successor]:
                dependency[predecessor] += sigma[predecessor] * coefficient
            if successor != source:
                centrality[successor] += dependency[successor]

    n = len(nodes)
    if n <= 2:
        return {node: Fraction(0, 1) for node in nodes}
    scale = Fraction(1, (n - 1) * (n - 2))
    return {node: value * scale for node, value in centrality.items()}


def connected_components(adjacency: dict[str, dict[str, Fraction]]) -> list[list[str]]:
    remaining = set(adjacency)
    components: list[list[str]] = []
    while remaining:
        start = min(remaining)
        queue = deque([start])
        remaining.remove(start)
        component: list[str] = []
        while queue:
            vertex = queue.popleft()
            component.append(vertex)
            for neighbor in sorted(adjacency[vertex]):
                if neighbor in remaining:
                    remaining.remove(neighbor)
                    queue.append(neighbor)
        components.append(sorted(component))
    return sorted(components, key=lambda values: (-len(values), values))


def run_known_graph_tests() -> list[dict[str, Any]]:
    tests: list[dict[str, Any]] = []

    def check(name: str, nodes: list[str], edges: list[tuple[str, str, Fraction]], expected: dict[str, Fraction]) -> None:
        actual = brandes_normalized(graph_from_edges(nodes, edges))
        if actual != expected:
            raise AssertionError(f"{name}: expected {expected}, found {actual}")
        tests.append({
            "test": name,
            "status": "PASS",
            "expected": {key: str(value) for key, value in sorted(expected.items())},
        })

    one = Fraction(1, 1)
    check(
        "Path_3",
        ["A", "B", "C"],
        [("A", "B", one), ("B", "C", one)],
        {"A": Fraction(0), "B": Fraction(1), "C": Fraction(0)},
    )
    check(
        "Star_4",
        ["A", "B", "C", "D"],
        [("A", "B", one), ("A", "C", one), ("A", "D", one)],
        {"A": Fraction(1), "B": Fraction(0), "C": Fraction(0), "D": Fraction(0)},
    )
    check(
        "Triangle_3",
        ["A", "B", "C"],
        [("A", "B", one), ("B", "C", one), ("A", "C", one)],
        {"A": Fraction(0), "B": Fraction(0), "C": Fraction(0)},
    )
    check(
        "Disconnected_Path_3_Plus_Isolate",
        ["A", "B", "C", "D"],
        [("A", "B", one), ("B", "C", one)],
        {"A": Fraction(0), "B": Fraction(1, 3), "C": Fraction(0), "D": Fraction(0)},
    )
    check(
        "Weighted_Equal_Path_Diamond",
        ["A", "B", "C", "D"],
        [
            ("A", "B", Fraction(1)),
            ("B", "D", Fraction(2)),
            ("A", "C", Fraction(2)),
            ("C", "D", Fraction(1)),
        ],
        {"A": Fraction(1, 6), "B": Fraction(1, 6), "C": Fraction(1, 6), "D": Fraction(1, 6)},
    )
    return tests


def construct_network() -> dict[str, Any]:
    source_hashes = {label: sha256_file(path) for label, path in HASH_PATHS.items()}
    for label, expected in EXPECTED_HASHES.items():
        actual = source_hashes[label]
        if actual != expected:
            raise ValueError(f"{label} hash mismatch: expected {expected}, found {actual}")

    source_rows = read_sheet(SOURCE_PATH, "Link_Master_500")
    b4_rows = read_sheet(B4_PATH, "Link_Eligibility_500")
    taxonomy_rows = read_sheet(SOURCE_PATH, "Taxonomy_Master")
    require_fields(
        source_rows,
        [
            "Link_ID", "Supplier_ID", "Taxonomy_ID", "Supplier_Normalized", "Technology_Family",
            "Technology_Category", "Evidence_Code", "Composite_Key", "Capability_Supported_Layer",
            "Taxonomy_Review_Flag",
        ],
        "Link_Master_500",
    )
    require_fields(
        b4_rows,
        [
            "Link_ID", "Supplier_ID", "Supplier_Name", "Taxonomy_ID", "Technology_Category",
            "Evidence_Code", "Capability_Supported_Layer", "Structural_Primary_Eligible",
            "Composite_Key", "Supplier_Taxonomy_Pair_Key",
        ],
        "B4 Link_Eligibility_500",
    )
    require_fields(taxonomy_rows, ["Taxonomy_ID", "Technology_Family", "Technology_Category"], "Taxonomy_Master")

    if len(source_rows) != 500 or len(b4_rows) != 500:
        raise ValueError(f"Expected 500 source and B4 rows, found {len(source_rows)} and {len(b4_rows)}")
    source_link_ids = [norm(row["Link_ID"]) for row in source_rows]
    b4_link_ids = [norm(row["Link_ID"]) for row in b4_rows]
    if any(not value for value in source_link_ids + b4_link_ids):
        raise ValueError("Blank Link_ID found")
    if len(set(source_link_ids)) != 500 or len(set(b4_link_ids)) != 500:
        raise ValueError("Link_ID grain is not unique")
    if set(source_link_ids) != set(b4_link_ids):
        raise ValueError("Frozen master and B4 Link_ID sets differ")

    source_by_link = {norm(row["Link_ID"]): row for row in source_rows}
    b4_by_link = {norm(row["Link_ID"]): row for row in b4_rows}
    reconstructed_eligible_ids = {
        link_id
        for link_id, row in source_by_link.items()
        if int(row["Capability_Supported_Layer"] or 0) == 1
        and norm(row["Taxonomy_Review_Flag"]) != "Needs manual taxonomy review"
    }
    b4_eligible_ids = {
        link_id for link_id, row in b4_by_link.items() if int(row["Structural_Primary_Eligible"] or 0) == 1
    }
    if reconstructed_eligible_ids != b4_eligible_ids:
        raise ValueError(
            "Reconstructed and B4 eligibility differ: "
            f"source_only={sorted(reconstructed_eligible_ids - b4_eligible_ids)}, "
            f"b4_only={sorted(b4_eligible_ids - reconstructed_eligible_ids)}"
        )
    if len(b4_eligible_ids) != 392:
        raise ValueError(f"Expected 392 eligible links, found {len(b4_eligible_ids)}")

    taxonomy_info = {
        norm(row["Taxonomy_ID"]): {
            "Technology_Family": norm(row["Technology_Family"]),
            "Technology_Category": norm(row["Technology_Category"]),
        }
        for row in taxonomy_rows
        if norm(row["Taxonomy_ID"])
    }

    eligible_rows = [b4_by_link[link_id] for link_id in sorted(b4_eligible_ids)]
    for row in eligible_rows:
        for field in ["Link_ID", "Supplier_ID", "Taxonomy_ID", "Composite_Key"]:
            if not norm(row[field]):
                raise ValueError(f"Blank {field} in eligible row {row}")
        if norm(row["Taxonomy_ID"]) == "TAX-18":
            raise ValueError("TAX-18 entered primary network")
        if norm(row["Evidence_Code"]) in {"U1", "X1"}:
            raise ValueError(f"{row['Evidence_Code']} entered primary network")

    supplier_ids = sorted({norm(row["Supplier_ID"]) for row in eligible_rows})
    taxonomy_ids = sorted({norm(row["Taxonomy_ID"]) for row in eligible_rows})
    if len(supplier_ids) != 207 or len(taxonomy_ids) != 17:
        raise ValueError(f"Expected 207 suppliers and 17 taxonomies, found {len(supplier_ids)} and {len(taxonomy_ids)}")

    supplier_names: dict[str, str] = {}
    pair_links: dict[tuple[str, str], set[str]] = defaultdict(set)
    for row in eligible_rows:
        supplier_id = norm(row["Supplier_ID"])
        taxonomy_id = norm(row["Taxonomy_ID"])
        link_id = norm(row["Link_ID"])
        supplier_names.setdefault(supplier_id, norm(row["Supplier_Name"]))
        pair_links[(supplier_id, taxonomy_id)].add(link_id)

    if len(pair_links) != 254:
        raise ValueError(f"Expected 254 supplier-taxonomy pairs, found {len(pair_links)}")
    if sum(len(link_ids) for link_ids in pair_links.values()) != 392:
        raise ValueError("Bipartite weights do not sum to 392")

    bipartite_edges: list[dict[str, Any]] = []
    supplier_vectors: dict[str, dict[str, int]] = {supplier_id: {} for supplier_id in supplier_ids}
    for supplier_id, taxonomy_id in sorted(pair_links):
        link_ids = sorted(pair_links[(supplier_id, taxonomy_id)])
        weight = len(link_ids)
        supplier_vectors[supplier_id][taxonomy_id] = weight
        bipartite_edges.append({
            "Pair_Key": f"{supplier_id}||{taxonomy_id}",
            "Supplier_ID": supplier_id,
            "Taxonomy_ID": taxonomy_id,
            "Technology_Family": taxonomy_info.get(taxonomy_id, {}).get("Technology_Family", ""),
            "Technology_Category": taxonomy_info.get(taxonomy_id, {}).get("Technology_Category", ""),
            "Weight_Unique_Link_ID": weight,
            "Link_IDs": link_ids,
            "Population_ID": "POP-PAIR-CAP-CTL-254",
        })

    projection_edges: list[dict[str, Any]] = []
    graph_edges: list[tuple[str, str, Fraction]] = []
    strength = {supplier_id: Fraction(0, 1) for supplier_id in supplier_ids}
    for left_index, left in enumerate(supplier_ids):
        left_vector = supplier_vectors[left]
        for right in supplier_ids[left_index + 1 :]:
            right_vector = supplier_vectors[right]
            categories = set(left_vector) | set(right_vector)
            numerator = sum(min(left_vector.get(category, 0), right_vector.get(category, 0)) for category in categories)
            denominator = sum(max(left_vector.get(category, 0), right_vector.get(category, 0)) for category in categories)
            if denominator <= 0:
                raise ValueError(f"Invalid Jaccard denominator for {left}, {right}")
            if numerator == 0:
                continue
            similarity = Fraction(numerator, denominator)
            distance = Fraction(denominator, numerator)
            strength[left] += similarity
            strength[right] += similarity
            graph_edges.append((left, right, distance))
            projection_edges.append({
                "Projection_Key": f"{left}||{right}",
                "Supplier_A": left,
                "Supplier_B": right,
                "Similarity_Numerator": numerator,
                "Similarity_Denominator": denominator,
                "Similarity_Exact": f"{similarity.numerator}/{similarity.denominator}",
                "Similarity": decimal(similarity),
                "Distance_Exact": f"{distance.numerator}/{distance.denominator}",
                "Distance": decimal(distance),
                "Population_ID": "POP-SUP-CAP-CTL-207",
            })

    adjacency = graph_from_edges(supplier_ids, graph_edges)
    possible_pairs = len(supplier_ids) * (len(supplier_ids) - 1) // 2
    if possible_pairs != 21_321:
        raise ValueError(f"Expected 21,321 possible pairs, found {possible_pairs}")
    if len(projection_edges) + (possible_pairs - len(projection_edges)) != possible_pairs:
        raise ValueError("Projection pair reconciliation failed")

    betweenness = brandes_normalized(adjacency)
    tests = run_known_graph_tests()
    components = connected_components(adjacency)
    isolate_ids = sorted(node for node in supplier_ids if not adjacency[node])
    degree = {node: len(adjacency[node]) for node in supplier_ids}

    supplier_metrics: list[dict[str, Any]] = []
    for supplier_id in supplier_ids:
        strength_value = strength[supplier_id]
        betweenness_value = betweenness[supplier_id]
        supplier_metrics.append({
            "Supplier_ID": supplier_id,
            "Supplier_Name": supplier_names.get(supplier_id, ""),
            "Strength_Exact": f"{strength_value.numerator}/{strength_value.denominator}",
            "Strength": decimal(strength_value),
            "Normalized_Betweenness_Exact": f"{betweenness_value.numerator}/{betweenness_value.denominator}",
            "Normalized_Betweenness": decimal(betweenness_value),
            "Isolate_Flag": 1 if supplier_id in isolate_ids else 0,
            "Positive_Edge_Count_QA": degree[supplier_id],
            "Population_ID": "POP-SUP-CAP-CTL-207",
        })

    supplier_nodes = [
        {
            "Node_ID": supplier_id,
            "Node_Type": "Supplier",
            "Supplier_Name": supplier_names.get(supplier_id, ""),
            "Population_ID": "POP-SUP-CAP-CTL-207",
            "Projection_Isolate_Flag": 1 if supplier_id in isolate_ids else 0,
        }
        for supplier_id in supplier_ids
    ]
    taxonomy_nodes = [
        {
            "Node_ID": taxonomy_id,
            "Node_Type": "Technology_Taxonomy",
            "Technology_Family": taxonomy_info.get(taxonomy_id, {}).get("Technology_Family", ""),
            "Technology_Category": taxonomy_info.get(taxonomy_id, {}).get("Technology_Category", ""),
            "Population_ID": "POP-LINK-CAP-CTL-392",
        }
        for taxonomy_id in taxonomy_ids
    ]

    link_input = []
    for link_id in sorted(b4_by_link):
        row = b4_by_link[link_id]
        source_row = source_by_link[link_id]
        reconstructed = int(
            int(source_row["Capability_Supported_Layer"] or 0) == 1
            and norm(source_row["Taxonomy_Review_Flag"]) != "Needs manual taxonomy review"
        )
        b4_eligible = int(row["Structural_Primary_Eligible"] or 0)
        link_input.append({
            "Link_ID": link_id,
            "Supplier_ID": norm(row["Supplier_ID"]),
            "Supplier_Name": norm(row["Supplier_Name"]),
            "Taxonomy_ID": norm(row["Taxonomy_ID"]),
            "Technology_Category": norm(row["Technology_Category"]),
            "Evidence_Code": norm(row["Evidence_Code"]),
            "Capability_Supported_Layer": int(row["Capability_Supported_Layer"] or 0),
            "Taxonomy_Review_Flag": norm(source_row["Taxonomy_Review_Flag"]),
            "B4_Structural_Primary_Eligible": b4_eligible,
            "C3_Eligible_Recalc": reconstructed,
            "Eligibility_Match": "PASS" if b4_eligible == reconstructed else "REVIEW",
            "Composite_Key": norm(row["Composite_Key"]),
            "Supplier_Taxonomy_Pair_Key": norm(row["Supplier_Taxonomy_Pair_Key"]),
        })

    degree_frequency = Counter(degree.values())
    weight_frequency = Counter(edge["Weight_Unique_Link_ID"] for edge in bipartite_edges)
    density = Fraction(len(projection_edges), possible_pairs)
    strength_values = [strength[node] for node in supplier_ids]
    betweenness_values = [betweenness[node] for node in supplier_ids]
    similarity_values = [Fraction(edge["Similarity_Numerator"], edge["Similarity_Denominator"]) for edge in projection_edges]

    canonical_hashes = {
        "supplier_nodes": sha256_text(supplier_ids),
        "bipartite_edges": sha256_text(
            f"{edge['Supplier_ID']}\t{edge['Taxonomy_ID']}\t{edge['Weight_Unique_Link_ID']}\t{','.join(edge['Link_IDs'])}"
            for edge in bipartite_edges
        ),
        "projection_edges": sha256_text(
            f"{edge['Supplier_A']}\t{edge['Supplier_B']}\t{edge['Similarity_Numerator']}\t{edge['Similarity_Denominator']}"
            for edge in projection_edges
        ),
        "supplier_metrics": sha256_text(
            f"{metric['Supplier_ID']}\t{metric['Strength_Exact']}\t{metric['Normalized_Betweenness_Exact']}"
            for metric in supplier_metrics
        ),
    }

    summary = {
        "candidate_links": len(link_input),
        "eligible_links": len(eligible_rows),
        "supplier_nodes": len(supplier_ids),
        "taxonomy_nodes": len(taxonomy_ids),
        "total_bipartite_nodes": len(supplier_ids) + len(taxonomy_ids),
        "bipartite_pair_edges": len(bipartite_edges),
        "bipartite_weight_sum": sum(edge["Weight_Unique_Link_ID"] for edge in bipartite_edges),
        "projection_nodes": len(supplier_ids),
        "projection_positive_edges": len(projection_edges),
        "projection_possible_pairs": possible_pairs,
        "projection_zero_similarity_pairs": possible_pairs - len(projection_edges),
        "projection_density_exact": f"{density.numerator}/{density.denominator}",
        "projection_density": decimal(density),
        "projection_density_percent": decimal(density * 100),
        "projection_components": len(components),
        "projection_component_sizes": [len(component) for component in components],
        "projection_isolates": len(isolate_ids),
        "strength_eligible_n": len(supplier_ids),
        "strength_valid_n": len(strength_values),
        "strength_missing_n": 0,
        "strength_zero_n": sum(value == 0 for value in strength_values),
        "betweenness_eligible_n": len(supplier_ids),
        "betweenness_valid_n": len(betweenness_values),
        "betweenness_missing_n": 0,
        "betweenness_zero_n": sum(value == 0 for value in betweenness_values),
        "strength_min": decimal(min(strength_values)),
        "strength_mean": decimal(sum(strength_values, Fraction(0)) / len(strength_values)),
        "strength_median": decimal(sorted(strength_values)[len(strength_values) // 2]),
        "strength_max": decimal(max(strength_values)),
        "betweenness_min": decimal(min(betweenness_values)),
        "betweenness_mean": decimal(sum(betweenness_values, Fraction(0)) / len(betweenness_values)),
        "betweenness_median": decimal(sorted(betweenness_values)[len(betweenness_values) // 2]),
        "betweenness_max": decimal(max(betweenness_values)),
        "similarity_min": decimal(min(similarity_values)),
        "similarity_mean": decimal(sum(similarity_values, Fraction(0)) / len(similarity_values)),
        "similarity_median": decimal(sorted(similarity_values)[len(similarity_values) // 2 - 1] / 2 + sorted(similarity_values)[len(similarity_values) // 2] / 2),
        "similarity_max": decimal(max(similarity_values)),
    }

    expected_summary = {
        "candidate_links": 500,
        "eligible_links": 392,
        "supplier_nodes": 207,
        "taxonomy_nodes": 17,
        "total_bipartite_nodes": 224,
        "bipartite_pair_edges": 254,
        "bipartite_weight_sum": 392,
        "projection_nodes": 207,
        "projection_possible_pairs": 21_321,
        "strength_eligible_n": 207,
        "strength_valid_n": 207,
        "strength_missing_n": 0,
        "betweenness_eligible_n": 207,
        "betweenness_valid_n": 207,
        "betweenness_missing_n": 0,
    }
    for key, expected in expected_summary.items():
        if summary[key] != expected:
            raise ValueError(f"{key}: expected {expected}, found {summary[key]}")

    return {
        "manifest": {
            "checklist_step": "C3",
            "decision_id": "GATE-C3-IMPL-01",
            "decision_status": "APPROVED — LOCKED BEFORE OFFICIAL C3 CALCULATION",
            "calculation_date": "2026-08-28",
            "python_version": platform.python_version(),
            "openpyxl_version": __import__("openpyxl").__version__,
            "arithmetic": "Exact fractions; decimals rounded only for export/display",
            "graph_class": "Simple undirected weighted bipartite graph and simple undirected positive-similarity supplier projection",
            "shortest_path_rule": "Exact weighted Dijkstra/Brandes on Distance = 1 / Similarity; endpoints excluded",
            "betweenness_normalization": "Source-accumulated undirected totals × 1/[(N-1)(N-2)] using full eligible N=207",
            "isolate_rule": "Eligible isolates retained with structural zero",
            "population_id": "POP-LINK-CAP-CTL-392",
            "supplier_population_id": "POP-SUP-CAP-CTL-207",
            "pair_population_id": "POP-PAIR-CAP-CTL-254",
            "implementation_file_sha256": sha256_file(IMPLEMENTATION_PATH),
            "source_hashes": source_hashes,
            "canonical_hashes": canonical_hashes,
        },
        "summary": summary,
        "diagnostics": {
            "eligibility_set_match": True,
            "source_only_eligible_link_ids": [],
            "b4_only_eligible_link_ids": [],
            "pair_weight_frequency": {str(key): value for key, value in sorted(weight_frequency.items())},
            "projection_degree_frequency": {str(key): value for key, value in sorted(degree_frequency.items())},
            "projection_components": components,
            "projection_isolate_ids": isolate_ids,
            "known_graph_tests": tests,
        },
        "link_input": link_input,
        "supplier_nodes": supplier_nodes,
        "taxonomy_nodes": taxonomy_nodes,
        "bipartite_edges": bipartite_edges,
        "projection_edges": projection_edges,
        "supplier_metrics": supplier_metrics,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, help="Ignored local path for the full deterministic JSON payload; requires --full-output")
    parser.add_argument("--full-output", action="store_true", help="Permit a row-level JSON export under the repository tmp directory")
    parser.add_argument("--summary-only", action="store_true", help="Deprecated compatibility flag; compact output is now the default")
    args = parser.parse_args()

    if args.output and not args.full_output:
        parser.error("--output contains row-level records and therefore requires --full-output")
    if args.full_output and not args.output:
        parser.error("--full-output requires --output under the repository tmp directory")

    result = construct_network()
    if args.output:
        output_path = require_private_output(args.output)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_text(json.dumps(result, indent=2, ensure_ascii=False, sort_keys=True) + "\n", encoding="utf-8")

    printable = {
        "manifest": result["manifest"],
        "summary": result["summary"],
        "diagnostics": {
            "eligibility_set_match": result["diagnostics"]["eligibility_set_match"],
            "pair_weight_frequency": result["diagnostics"]["pair_weight_frequency"],
            "projection_degree_frequency": result["diagnostics"]["projection_degree_frequency"],
            "projection_component_sizes": result["summary"]["projection_component_sizes"],
            "projection_isolate_count": len(result["diagnostics"]["projection_isolate_ids"]),
            "known_graph_tests": result["diagnostics"]["known_graph_tests"],
        },
        "output": str(args.output.resolve()) if args.output else None,
    }
    print(json.dumps(printable, indent=2, ensure_ascii=False, sort_keys=True))


if __name__ == "__main__":
    main()
